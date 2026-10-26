#!/usr/bin/env python3
"""
Generates the Introduction section of the T-DPC-RWA study as:
  Introduction.docx      Word file (Times New Roman 12 pt) that opens cleanly in Word and LibreOffice
  introduction.md        the same text as Markdown
  fig1_workflow.mmd      Figure 1 as Mermaid source
  fig1_workflow.drawio   Figure 1 as draw.io XML (open at app.diagrams.net)

Requires: pip install python-docx
Run:      python make_introduction.py [output_dir]

Citations are written as [@key] or [@key1,@key2] and are numbered automatically
in order of first appearance (IEEE style), so text, table, caption and reference
list always agree.
"""
import os
import re
import sys
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
FONT = "Times New Roman"
BODY_PT = 12
TABLE_PT = 10      # table and figure text; set to 12 if a uniform 12 pt is required

# ============================================================================ references
REFS = {
    "bb84": 'C. H. Bennett and G. Brassard, "Quantum cryptography: Public key distribution and coin tossing," in *Proc. IEEE Int. Conf. Computers, Systems and Signal Processing*, Bangalore, India, 1984, pp. 175–179.',
    "scarani": 'V. Scarani, H. Bechmann-Pasquinucci, N. J. Cerf, M. Dušek, N. Lütkenhaus, and M. Peev, "The security of practical quantum key distribution," *Rev. Mod. Phys.*, vol. 81, no. 3, p. 1301, 2009.',
    "decoy": 'H.-K. Lo, X. Ma, and K. Chen, "Decoy state quantum key distribution," *Phys. Rev. Lett.*, vol. 94, no. 23, p. 230504, 2005.',
    "caosurvey": 'Y. Cao, Y. Zhao, Q. Wang, J. Zhang, S. X. Ng, and L. Hanzo, "The evolution of quantum key distribution networks: On the road to the Qinternet," *IEEE Commun. Surveys Tuts.*, vol. 24, pp. 839–894, 2022.',
    "sdqaas": 'Y. Cao, Y. Zhao, J. Wang, X. Yu, Z. Ma, and J. Zhang, "SDQaaS: Software defined networking for quantum key distribution as a service," *Opt. Express*, vol. 27, no. 5, pp. 6892–6909, 2019.',
    "wangsdn": 'H. Wang, Y. Zhao, and A. Nag, "Quantum-key-distribution (QKD) networks enabled by software-defined networks (SDN)," *Appl. Sci.*, vol. 9, no. 10, p. 2081, 2019.',
    "chapuran": 'T. E. Chapuran *et al.*, "Optical networking for quantum key distribution and quantum communications," *New J. Phys.*, vol. 11, p. 105001, 2009.',
    "peters": 'N. A. Peters *et al.*, "Dense wavelength multiplexing of 1550 nm QKD with strong classical channels in reconfigurable networking environments," *New J. Phys.*, vol. 11, p. 045012, 2009.',
    "eraerds": 'P. Eraerds, N. Walenta, M. Legré, N. Gisin, and H. Zbinden, "Quantum key distribution and 1 Gbps data encryption over a single fibre," *New J. Phys.*, vol. 12, p. 063027, 2010.',
    "patel": 'K. A. Patel *et al.*, "Coexistence of high-bit-rate quantum key distribution and data on optical fiber," *Phys. Rev. X*, vol. 2, no. 4, p. 041010, 2012.',
    "dasilva": 'T. Ferreira da Silva, G. B. Xavier, G. P. Temporão, and J. P. von der Weid, "Impact of Raman scattered noise from multiple telecom channels on fiber-optic quantum key distribution systems," *J. Lightw. Technol.*, vol. 32, pp. 2332–2339, 2014.',
    "chlamtac": 'I. Chlamtac, A. Ganz, and G. Karmi, "Lightpath communications: An approach to high bandwidth optical WAN\'s," *IEEE Trans. Commun.*, vol. 40, pp. 1171–1182, 1992.',
    "zang": 'H. Zang, J. P. Jue, and B. Mukherjee, "A review of routing and wavelength assignment approaches for wavelength-routed optical WDM networks," *Opt. Netw. Mag.*, vol. 1, pp. 47–60, 2000.',
    "cao2017": 'Y. Cao, Y. Zhao, X. Yu, and Y. Wu, "Resource assignment strategy in optical networks integrated with quantum key distribution," *J. Opt. Commun. Netw.*, vol. 9, no. 11, pp. 995–1004, 2017.',
    "zhao2018": 'Y. Zhao *et al.*, "Resource allocation in optical networks secured by quantum key distribution," *IEEE Commun. Mag.*, vol. 56, pp. 130–137, 2018.',
    "dong": 'K. Dong, Y. Zhao, X. Yu, A. Nag, and J. Zhang, "Auxiliary graph based routing, wavelength, and time-slot assignment in metro quantum optical networks with a novel node structure," *Opt. Express*, vol. 28, no. 5, pp. 5936–5952, 2020.',
    "zhangq": 'Q. Zhang, O. Ayoub, A. Gatto, J. Wu, F. Musumeci, and M. Tornatore, "Routing, channel, key-rate, and time-slot assignment for QKD in optical networks," *IEEE Trans. Netw. Serv. Manag.*, early access, 2023.',
    "rawc": 'T. Zhao, X. Fan, B. Dong, Q. Niu, and B. Guo, "A resource-adaptive routing scheme with wavelength conflicts in quantum key distribution optical networks," *Entropy*, vol. 25, no. 5, p. 732, 2023.',
    "bahrani": 'S. Bahrani, M. Razavi, and J. A. Salehi, "Wavelength assignment in hybrid quantum-classical networks," *Sci. Rep.*, vol. 8, p. 3456, 2018.',
    "moghaddam": 'E. E. Moghaddam, H. Beyranvand, and J. A. Salehi, "Resource allocation in space division multiplexed elastic optical networks secured with quantum key distribution," *IEEE J. Sel. Areas Commun.*, vol. 39, no. 9, pp. 2688–2700, 2021.',
    "kiktenko": 'E. O. Kiktenko, A. Tayduganov, and A. K. Fedorov, "Routing algorithm within the multiple non-overlapping paths\' approach for quantum key distribution networks," *Entropy*, vol. 26, no. 12, p. 1102, 2024.',
    "lella": 'E. Lella and G. Schmid, "On the security of quantum key distribution networks," *Cryptography*, vol. 7, no. 4, p. 53, 2023.',
    "ruiz2025": 'L. Ruiz and J. C. Garcia-Escartin, "Routing and wavelength assignment in hybrid networks with classical and quantum signals," *IEEE J. Sel. Areas Commun.*, vol. 43, no. 2, pp. 412–421, 2025.',
    "ruiz2026": 'L. Ruiz and J. C. Garcia-Escartin, "Routing and wavelength assignment in quantum key distribution networks: Power control heuristics for quantum–classical multiplexing," *Comput. Commun.*, vol. 256, p. 108572, 2026.',
    "gabriel": 'K. R. Gabriel and R. R. Sokal, "A new statistical approach to geographic variation analysis," *Syst. Zool.*, vol. 18, no. 3, pp. 259–278, 1969.',
    "waxman": 'B. M. Waxman, "Routing of multipoint connections," *IEEE J. Sel. Areas Commun.*, vol. 6, no. 9, pp. 1617–1622, 1988.',
    "birand": 'B. Birand, H. Wang, K. Bergman, D. Kilper, T. Nandagopal, and G. Zussman, "Real-time power control for dynamic optical networks—Algorithms and experimentation," *IEEE J. Sel. Areas Commun.*, vol. 32, pp. 1615–1628, 2014.',
    "subramaniam": 'S. Subramaniam, M. Azizoglu, and A. K. Somani, "All-optical networks with sparse wavelength conversion," *IEEE/ACM Trans. Netw.*, vol. 4, no. 4, pp. 544–557, 1996.',
    "brandes": 'U. Brandes, "A faster algorithm for betweenness centrality," *J. Math. Sociol.*, vol. 25, no. 2, pp. 163–177, 2001.',
}

# ============================================================================ text
TITLE = "1. Introduction"

SECTIONS = [
    ("1.1 Background and Context", [
        "The security of most deployed public-key cryptography rests on the presumed difficulty of problems such as integer factorisation and discrete logarithms, an assumption that large-scale quantum computation would undermine. Quantum key distribution (QKD) offers a different foundation: two parties expand a shared secret over an insecure quantum channel and an authenticated classical channel, and any attempt to intercept the quantum signals leaves a measurable trace in the observed error rate. Since the first protocol was proposed by Bennett and Brassard [@bb84], the security of practical implementations with imperfect sources and detectors has been analysed in depth [@scarani], and the decoy-state technique of Lo, Ma and Chen made weak coherent laser sources compatible with rigorous security bounds [@decoy].",
        "As point-to-point QKD links matured, attention moved to networks in which many node pairs obtain keys on demand. Cao et al. trace this evolution from isolated testbeds towards a large-scale quantum-secured infrastructure [@caosurvey], and several groups have shown how a software-defined control plane can expose QKD resources as programmable network services [@sdqaas,@wangsdn]. Dedicated dark fibre for every quantum link is, however, rarely affordable outside trial deployments. A sustainable path is to carry quantum channels over the same wavelength-division-multiplexed (WDM) fibre plant that already transports classical traffic.",
        "Sharing a fibre between these two kinds of signal is physically demanding. A classical channel is launched at powers in the milliwatt range, whereas a quantum channel carries, on average, less than one photon per pulse. Filters suppress linear crosstalk effectively, but nonlinear processes in silica create new photons at the quantum wavelength that cannot be filtered away. Chapuran et al. and Peters et al. identified spontaneous Raman scattering as the dominant impairment when a quantum channel coexists with strong classical signals [@chapuran,@peters]. Eraerds et al. operated QKD alongside gigabit data on a single fibre [@eraerds], Patel et al. used temporal filtering to extend coexistence to 90 km [@patel], and Ferreira da Silva et al. quantified how Raman noise from many telecom channels accumulates in a quantum receiver [@dasilva]. A widely adopted mitigation places the quantum channel in the O-band near 1310 nm and the classical traffic in the C-band near 1550 nm, so that the quantum signal sees only the weak anti-Stokes tail of the Raman spectrum, at the price of higher fibre attenuation.",
        "These findings move part of the security problem into the network layer. Which classical lightpaths share fibre with a quantum lightpath, over what distance and at what launch power, is decided by routing and wavelength assignment (RWA). For classical optical networks RWA is known to be computationally hard [@chlamtac], and practical networks rely on heuristics such as shortest-path first-fit [@zang]. In a hybrid network, every routing decision for a classical channel also changes the noise floor of the quantum channels it overlaps, so RWA must respect a quantum signal-to-noise constraint in addition to wavelength availability.",
    ]),
    ("1.2 Literature Review and Related Work", [
        "A substantial body of work addresses resource allocation for QKD-secured optical networks. Cao et al. formulated an integer linear programme and companion heuristics that assign routes, wavelengths and time slots to the quantum, synchronisation and data channels of each QKD session [@cao2017], and Zhao et al. framed resource allocation in QKD-secured optical networks as a general architectural problem [@zhao2018]. Dong et al. introduced a node structure able to bypass intermediate nodes optically and an auxiliary-graph algorithm that exploits it [@dong]. Zhang et al. jointly optimised routing, channel, key-rate and time-slot assignment [@zhangq], while Zhao et al. adjusted link weights according to a wavelength-conflict metric to raise the acceptance of simultaneous QKD requests [@rawc].",
        "Other studies concentrate on how spectrum and space are shared. Bahrani et al. showed that interleaving quantum and classical sub-bands, rather than separating them into two contiguous blocks, can increase the total secret key rate in noise-limited regimes [@bahrani]. Moghaddam et al. extended the allocation problem to space-division-multiplexed elastic optical networks built on multi-core fibre [@moghaddam]. When end-to-end distances exceed what a single QKD link supports, keys are usually forwarded through trusted relay nodes; Kiktenko et al. reduced the trust placed in individual relays by spreading key material over non-overlapping paths [@kiktenko], and Lella and Schmid examined the security assumptions that relay-based networks impose [@lella].",
        "Closest to the present work are the physics-aware heuristics of Ruiz and Garcia-Escartin. Their first study ranked candidate routes for classical channels by the length of fibre they share with established quantum channels, leading to minimum-overlap and fully disjoint strategies [@ruiz2025]. A subsequent study added launch-power control: once a route is chosen, the classical transmitter power is lowered to the minimum that still satisfies the receiver signal-to-noise threshold, which reduces the Raman photons injected into co-propagating quantum channels [@ruiz2026]. Evaluated on random Gabriel [@gabriel] and Waxman [@waxman] topologies, power control lowered blocking and raised the quantum signal-to-noise ratio compared with fixed-power operation. Outside the quantum domain, Birand et al. demonstrated real-time power control in dynamic optical networks [@birand], and the classical RWA literature has long shown that a small number of well-placed wavelength converters recovers much of the benefit of full conversion [@subramaniam]. Table I compares these contributions.",
    ]),
    ("1.3 Research Gap and Limitations", [
        "The reviewed studies establish that resource allocation must account for the physical layer when quantum and classical signals share fibre. Several technical limitations, however, remain open.",
        "**Static demand models.** The power-controlled heuristics of [@ruiz2026] are evaluated on demand sets known in advance and provisioned sequentially, without departures. In operational networks, connections arrive and leave stochastically, so the set of classical interferers seen by a quantum channel changes throughout its lifetime. The authors themselves identify dynamic traffic assignment as an open question [@ruiz2026].",
        "**One-shot launch-power assignment.** Classical launch power is computed once, from the path length at set-up, and is never revisited. It therefore cannot react to new interferers, to transient bursts of classical traffic, or to slow drifts such as connector degradation, which may push a quantum channel below its viability threshold or a classical receiver below its sensitivity in mid-session.",
        "**Strict wavelength continuity.** The physics-aware heuristics in [@ruiz2025,@ruiz2026] assume networks without wavelength converters. Conversion below the single-photon level is impractical, but converting classical channels alone would not touch the quantum path, and the potential of sparse classical conversion [@subramaniam] has not been examined in hybrid networks.",
        "**Bounded single-segment reach.** The topologies in [@ruiz2026] are generated so that no path exceeds 60 km. This is a deliberate simulation choice rather than a physical limit of QKD, but it leaves open how physics-aware RWA should serve node pairs whose shortest path is too long for a single quantum segment. Trusted relays address reach in general [@caosurvey,@kiktenko], yet they have not been combined with Raman-aware routing and power control.",
        "**Open-loop control.** Software-defined QKD architectures [@sdqaas,@wangsdn] manage keys and services but do not close a control loop on physical-layer measurements such as the quantum signal-to-noise ratio. No existing scheme detects a quantum channel drifting towards outage and corrects it while the session is active.",
    ]),
    ("1.4 Proposed Contributions", [
        "This study develops and evaluates a Telemetry-Driven Dynamic Power Control and RWA framework (T-DPC-RWA) for hybrid quantum–classical networks under stochastic online traffic with Poisson arrivals and exponentially distributed holding times. Its objective is to extend physics-aware RWA from static provisioning to continuous operation, while keeping noise mitigation, spectrum relief and reach extension as separate mechanisms whose effects can be measured independently. The main contributions are:",
        "- A time-variant, segment-aware model of anti-Stokes Raman noise that yields closed-form sensitivities of each quantum channel to the launch power of each classical segment.",
        "- A closed-loop control algorithm that polls the quantum and classical signal-to-noise ratios at a fixed interval and, when a quantum channel enters a guard band, applies sensitivity-ordered power throttling and escalates to make-before-break re-grooming of the dominant interferer when throttling is insufficient, with hysteresis-based power recovery.",
        "- Sparse optical-electrical-optical wavelength conversion for classical channels only, placed at nodes of high betweenness centrality [@brandes] and exploited through a layered auxiliary graph, so that no active device enters the quantum path.",
        "- Trusted-relay segmentation that serves long-distance quantum requests through the relay chain with the highest bottleneck secret key rate under the decoy-state BB84 model [@decoy]; relays terminate and re-establish keys and are explicitly distinct from quantum repeaters.",
        "- A reproducible discrete-event evaluation on Gabriel and Waxman topologies that uses common random numbers, a factorial ablation of the three mechanisms and a decomposition of blocking into spectrum, noise and reach causes.",
        "Fig. 1 summarises the resulting workflow. The remainder of the paper presents the proposed architecture, the impairment model and algorithms, the simulation methodology and results, and concluding remarks.",
    ]),
]

# Table I: (Author(s) & Year [Citation], Methodology/Approach, Key findings/strengths, Technical limitations/research gaps)
TABLE = [
    ("Chapuran et al., 2009 [@chapuran]", "Experimental O-band QKD sharing fibre with amplified C-band data channels; Raman noise measurement and modelling.",
     "Identified spontaneous anti-Stokes Raman scattering as the dominant coexistence impairment.", "Link-level study; no routing or network-wide resource allocation."),
    ("Eraerds et al., 2010 [@eraerds]", "QKD and 1 Gbps encrypted data multiplexed on one fibre with DWDM classical channels.",
     "Showed practical single-fibre operation of QKD with its own classical channels.", "Fixed point-to-point link; no dynamic multi-user network."),
    ("Patel et al., 2012 [@patel]", "Temporal filtering of Raman noise photons in high-bit-rate QKD.",
     "Coexistence with bidirectional gigabit data over fibres up to 90 km.", "Hardware-level mitigation; does not address routing or power allocation."),
    ("Birand et al., 2014 [@birand]", "Real-time power control for dynamic classical optical networks, with experimental validation.",
     "Demonstrates feasible closed-loop power adjustment as channels are added and removed.", "Classical networks only; quantum signal-to-noise constraints not considered."),
    ("Cao et al., 2017 [@cao2017]", "Integer linear programming and heuristics for routing, wavelength and time-slot assignment in QKD-enabled networks.",
     "Near-optimal resource use for known traffic matrices.", "Static demands; Raman cross-talk between channels not modelled."),
    ("Bahrani et al., 2018 [@bahrani]", "Optimisation of wavelength positions for quantum and classical channels in DWDM links.",
     "Interleaved quantum and classical bands raise total secret key rate in noise-limited regimes.", "Single-link spectral design; no path selection or temporal dynamics."),
    ("Dong et al., 2020 [@dong]", "Auxiliary-graph RWTA with a node structure supporting optical bypass and trusted relaying.",
     "Bypassing intermediate nodes within QKD reach saves relay keys and resources.", "Physical-layer noise from co-propagating classical channels not considered."),
    ("Moghaddam et al., 2021 [@moghaddam]", "Resource allocation for QKD in space-division-multiplexed elastic optical networks.",
     "Extends QKD resource allocation to multi-core fibre infrastructures.", "Requires multi-core fibre; launch-power control not addressed."),
    ("Zhang et al., 2023 [@zhangq]", "Joint routing, channel, key-rate and time-slot assignment for QKD in optical networks.",
     "Unified treatment of key-rate demands within RWA.", "Raman-aware power control and closed-loop supervision absent."),
    ("Zhao et al., 2023 [@rawc]", "Resource-adaptive routing using a wavelength-conflict degree to update link weights.",
     "Up to 30% higher acceptance of simultaneous QKD requests.", "Spectral contention only; nonlinear noise and physical drift not modelled."),
    ("Kiktenko et al., 2024 [@kiktenko]", "Key routing over multiple non-overlapping paths in trusted-relay QKD networks.",
     "Security holds unless every path contains a compromised relay.", "Assumes key-generation capacities as given; physical-layer coexistence not modelled."),
    ("Ruiz and Garcia-Escartin, 2025 [@ruiz2025]", "Overlap-minimising and disjoint routing heuristics for classical channels in hybrid networks.",
     "Shorter shared spans reduce Raman noise using software-only rules.", "Fixed launch power; detours lengthen classical paths; static demands."),
    ("Ruiz and Garcia-Escartin, 2026 [@ruiz2026]", "Minimum-power launch control combined with physics-aware RWA heuristics on random geometric graphs.",
     "Lower blocking and higher quantum signal-to-noise ratio than fixed-power operation.", "Static sequential demands, one-shot power setting, no wavelength conversion, paths bounded at 60 km."),
]

FIG_CAPTION = ("Fig. 1. Conceptual workflow of the proposed T-DPC-RWA framework. Online provisioning (upper part) extends "
               "physics-aware RWA [@ruiz2025,@ruiz2026] with sparse classical wavelength conversion [@subramaniam] and "
               "trusted-relay segmentation; the telemetry loop (lower part) adds the closed-loop correction that static "
               "launch-power assignment lacks. Diagram by the authors.")

# Figure 1 as a grid: each row is ("full", box) or ("split", left, connector, right). Boxes: (text, fill) or None.
BLUE, GREEN, ORANGE, GREY = "DCE6F2", "E2EFDA", "FCE4D6", "F2F2F2"
FIG_ROWS = [
    ("full", ("Stochastic connection request r = (s, d, type): Poisson arrivals, exponential holding times", GREY)),
    ("split", ("▼ classical", None), "", ("▼ quantum", None)),
    ("split", ("Classical request: k candidate paths ranked by quantum-overlap cost; first-fit wavelength", BLUE), "",
     ("Quantum request: direct O-band lightpath on k shortest paths; test QSNR ≥ Γ_{Q} + ε", BLUE)),
    ("split", ("▼ no continuous free wavelength", None), "", ("▼ reach or QSNR test fails", None)),
    ("split", ("Layered auxiliary graph with conversion edges only at sparse O-E-O hubs (high betweenness)", BLUE), "",
     ("Trusted-relay chain s → r_{1} → … → d; keep the chain with the highest bottleneck SKR", BLUE)),
    ("split", ("▼", None), "", ("▼", None)),
    ("full", ("Admission: every active quantum channel keeps QSNR ≥ Γ_{Q} + ε; launch power P = P_{min}(1 + m_{0}); update network state", GREEN)),
    ("full", ("▼", None)),
    ("full", ("SDN telemetry poll every Δt: QSNR_{q}, SNR_{c,s}, excess loss η_{c,s}", ORANGE)),
    ("split", ("▼ QSNR < Γ_{Q} + ε", None), "", ("▼ QSNR ≥ Γ_{Q} + ε + h", None)),
    ("split", ("At-risk channel: compute noise reduction ΔN* and throttling headroom H", ORANGE), "",
     ("Hysteresis recovery: raise throttled powers by Δp towards P_{min}(1 + m_{0})", ORANGE)),
    ("split", ("Level 1 (H ≥ ΔN*): sensitivity-ordered power throttling", ORANGE), "→ if still at risk",
     ("Level 2: make-before-break re-grooming of the dominant interferer", ORANGE)),
    ("full", ("↺ repeat at the next telemetry interval Δt", None)),
]

# ============================================================================ citation numbering
ORDER: list[str] = []


def _scan(text):
    for group in re.findall(r"\[(@[^\]]+)\]", text):
        for key in re.findall(r"@(\w+)", group):
            if key not in REFS:
                raise KeyError(f"unknown citation key: {key}")
            if key not in ORDER:
                ORDER.append(key)


for _, paras in SECTIONS:
    for p in paras:
        _scan(p)
for row in TABLE:
    _scan(" ".join(row))
_scan(FIG_CAPTION)
NUM = {k: i + 1 for i, k in enumerate(ORDER)}


def _fmt_group(m):
    nums = sorted(NUM[k] for k in re.findall(r"@(\w+)", m.group(1)))
    out, i = [], 0
    while i < len(nums):          # compress runs of 3+ into ranges: [3]–[5]
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(f"[{nums[i]}]–[{nums[j]}]" if j - i >= 2 else ", ".join(f"[{n}]" for n in nums[i:j + 1]))
        i = j + 1
    return ", ".join(out)


def cite(text):
    return re.sub(r"\[(@[^\]]+)\]", _fmt_group, text)


# ============================================================================ docx helpers
TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{[^}]+\}|\^\{[^}]+\})")


def add_runs(par, text, size=BODY_PT, bold=False):
    """Inline markup: **bold**, *italic*, _{subscript}, ^{superscript}."""
    for part in TOKEN.split(text):
        if not part:
            continue
        b, it, sub, sup = bold, False, False, False
        if part.startswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("*"):
            part, it = part[1:-1], True
        elif part.startswith("_{"):
            part, sub = part[2:-1], True
        elif part.startswith("^{"):
            part, sup = part[2:-1], True
        run = par.add_run(part)
        run.bold, run.italic = b, it
        run.font.subscript, run.font.superscript = sub, sup
        run.font.size = Pt(size)
        run.font.name = FONT
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    return par


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_cell_borders(cell, size=8, color="404040", none=False):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if none:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        borders.append(el)
    tcpr.append(borders)


def set_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def keep_row_together(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def set_col_widths(table, widths_cm):
    table.autofit = False
    tbl = table._tbl
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def caption(doc, text, before=6, after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    add_runs(p, cite(text), size=TABLE_PT)
    return p


def table_par(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    add_runs(p, cite(text), size=TABLE_PT, bold=bold)


# ============================================================================ build .docx
def build_docx(path):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = FONT, Pt(size), True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    h = doc.add_heading(level=1)
    add_runs(h, TITLE, size=14, bold=True)

    for title, paras in SECTIONS:
        hh = doc.add_heading(level=2)
        add_runs(hh, title, size=12, bold=True)
        for text in paras:
            if text.startswith("- "):
                p = doc.add_paragraph(style="List Bullet")
                text = text[2:]
            else:
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0.75)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, cite(text))
        if title.startswith("1.2"):
            add_table(doc)
        if title.startswith("1.4"):
            add_figure(doc)

    hh = doc.add_heading(level=2)
    add_runs(hh, "References", size=12, bold=True)
    for key in ORDER:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.paragraph_format.space_after = Pt(3)
        add_runs(p, f"[{NUM[key]}]\t" + REFS[key], size=BODY_PT - 1)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
    doc.save(path)


def add_table(doc):
    caption(doc, "TABLE I\nComparative summary of related studies", before=10, after=4)
    headers = ["Author(s) & Year [Citation]", "Methodology / Approach", "Key Findings / Strengths",
               "Technical Limitations / Research Gaps"]
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(t.rows[0].cells, headers):
        table_par(cell, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, "D9D9D9")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_header(t.rows[0])
    for row in TABLE:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            table_par(cell, text)
        keep_row_together(t.rows[-1])
    set_col_widths(t, [3.4, 4.4, 4.2, 4.0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc):
    """Figure 1 drawn with native table cells: every box, label and arrow stays editable."""
    t = doc.add_table(rows=len(FIG_ROWS), cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [7.0, 2.0, 7.0]
    for r, spec in enumerate(FIG_ROWS):
        row = t.rows[r]
        keep_row_together(row)
        if spec[0] == "full":
            cell = row.cells[0].merge(row.cells[2])
            text, fill = spec[1]
            _fig_cell(cell, text, fill)
        else:
            (lt, lf), mid, (rt, rf) = spec[1], spec[2], spec[3]
            _fig_cell(row.cells[0], lt, lf)
            _fig_cell(row.cells[1], mid, None)
            _fig_cell(row.cells[2], rt, rf)
    set_col_widths(t, widths)
    caption(doc, FIG_CAPTION, before=6, after=10)


def _fig_cell(cell, text, fill):
    is_box = fill is not None
    set_cell_borders(cell, none=not is_box)
    if is_box:
        set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3 if is_box else 0)
    p.paragraph_format.space_after = Pt(3 if is_box else 0)
    p.paragraph_format.keep_with_next = True
    add_runs(p, text, size=TABLE_PT if is_box else TABLE_PT - 1, bold=False)
    if not is_box:
        for run in p.runs:
            run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)


# ============================================================================ markdown, mermaid, draw.io
def plain(text):
    """Markdown/diagram text: _{x} -> _x, ^{x} -> ^x."""
    return re.sub(r"[_^]\{([^}]+)\}", lambda m: ("_" if m.group(0)[0] == "_" else "^") + m.group(1), text)


def build_markdown(path):
    lines = [f"# {TITLE}", ""]
    for title, paras in SECTIONS:
        lines += [f"## {title}", ""]
        for text in paras:
            lines += [plain(cite(text)), ""]
        if title.startswith("1.2"):
            lines += ["**TABLE I. Comparative summary of related studies**", "",
                      "| Author(s) & Year [Citation] | Methodology / Approach | Key Findings / Strengths | Technical Limitations / Research Gaps |",
                      "|---|---|---|---|"]
            lines += ["| " + " | ".join(cite(c) for c in row) + " |" for row in TABLE]
            lines.append("")
        if title.startswith("1.4"):
            lines += ["*Figure 1: editable sources in `fig1_workflow.mmd` (Mermaid) and `fig1_workflow.drawio` (draw.io).*", "",
                      plain(cite(FIG_CAPTION)), ""]
    lines += ["## References", ""] + [f"[{NUM[k]}] {REFS[k]}" for k in ORDER] + [""]
    open(path, "w", encoding="utf-8").write("\n".join(lines))


FLOW_NODES = [  # id, label, kind (fill), x, y, w, h  (draw.io geometry in px)
    ("req", "Stochastic connection request r = (s, d, type)<br>Poisson arrivals, exponential holding times", GREY, 140, 0, 520, 60),
    ("cls", "Classical request:<br>k candidate paths ranked by quantum-overlap cost; first-fit wavelength", BLUE, 0, 110, 340, 70),
    ("qnt", "Quantum request:<br>direct O-band lightpath; test QSNR ≥ Γ_Q + ε", BLUE, 460, 110, 340, 70),
    ("aux", "Layered auxiliary graph<br>conversion edges only at sparse O-E-O hubs (high betweenness)", BLUE, 0, 240, 340, 70),
    ("rel", "Trusted-relay chain s → r1 → … → d<br>keep the chain with the highest bottleneck SKR", BLUE, 460, 240, 340, 70),
    ("adm", "Admission: all active quantum channels keep QSNR ≥ Γ_Q + ε<br>launch power P = P_min(1 + m0); update network state", GREEN, 140, 370, 520, 70),
    ("tel", "SDN telemetry poll every Δt<br>QSNR_q, SNR_c,s, excess loss η_c,s", ORANGE, 140, 490, 520, 60),
    ("risk", "At-risk channel:<br>compute noise reduction ΔN* and headroom H", ORANGE, 0, 610, 340, 60),
    ("rec", "Hysteresis recovery:<br>raise throttled powers by Δp towards P_min(1 + m0)", ORANGE, 460, 610, 340, 60),
    ("l1", "Level 1 (H ≥ ΔN*):<br>sensitivity-ordered power throttling", ORANGE, 0, 720, 340, 60),
    ("l2", "Level 2:<br>make-before-break re-grooming of the dominant interferer", ORANGE, 460, 720, 340, 60),
]
FLOW_EDGES = [
    ("req", "cls", "classical"), ("req", "qnt", "quantum"),
    ("cls", "aux", "no continuous free wavelength"), ("qnt", "rel", "reach or QSNR test fails"),
    ("cls", "adm", "path found"), ("qnt", "adm", "feasible"), ("aux", "adm", ""), ("rel", "adm", ""),
    ("adm", "tel", ""), ("tel", "risk", "QSNR < Γ_Q + ε"), ("tel", "rec", "QSNR ≥ Γ_Q + ε + h"),
    ("risk", "l1", "H ≥ ΔN*"), ("risk", "l2", "H < ΔN*"), ("l1", "l2", "still at risk"),
    ("l1", "tel", "next Δt"), ("l2", "tel", "next Δt"), ("rec", "tel", "next Δt"),
]


def build_mermaid(path):
    out = ["flowchart TD"]
    for nid, label, fill, *_ in FLOW_NODES:
        out.append(f'    {nid}["{label}"]')
    for a, b, lab in FLOW_EDGES:
        out.append(f'    {a} -->|"{lab}"| {b}' if lab else f"    {a} --> {b}")
    for nid, _, fill, *_ in FLOW_NODES:
        out.append(f"    style {nid} fill:#{fill},stroke:#404040")
    open(path, "w", encoding="utf-8").write("\n".join(out) + "\n")


def build_drawio(path):
    cells = ['<mxCell id="0"/>', '<mxCell id="1" parent="0"/>']
    for nid, label, fill, x, y, w, h in FLOW_NODES:
        style = f"rounded=1;whiteSpace=wrap;html=1;fillColor=#{fill};strokeColor=#404040;fontFamily=Times New Roman;fontSize=12;"
        cells.append(f'<mxCell id="{nid}" value="{escape(label, {chr(34): "&quot;"})}" style="{style}" vertex="1" parent="1">'
                     f'<mxGeometry x="{x}" y="{y}" width="{w}" height="{h}" as="geometry"/></mxCell>')
    for i, (a, b, lab) in enumerate(FLOW_EDGES):
        style = "edgeStyle=orthogonalEdgeStyle;rounded=0;html=1;endArrow=block;fontFamily=Times New Roman;fontSize=11;"
        cells.append(f'<mxCell id="e{i}" value="{escape(lab)}" style="{style}" edge="1" parent="1" source="{a}" target="{b}">'
                     f'<mxGeometry relative="1" as="geometry"/></mxCell>')
    xml = ('<mxfile host="app.diagrams.net"><diagram name="Fig1 T-DPC-RWA workflow">'
           '<mxGraphModel dx="900" dy="900" grid="1" gridSize="10" page="1" pageWidth="827" pageHeight="1169"><root>'
           + "".join(cells) + "</root></mxGraphModel></diagram></mxfile>\n")
    open(path, "w", encoding="utf-8").write(xml)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    build_docx(os.path.join(OUT, "Introduction.docx"))
    build_markdown(os.path.join(OUT, "introduction.md"))
    build_mermaid(os.path.join(OUT, "fig1_workflow.mmd"))
    build_drawio(os.path.join(OUT, "fig1_workflow.drawio"))
    print(f"Wrote Introduction.docx, introduction.md, fig1_workflow.mmd, fig1_workflow.drawio to {OUT}")
    print(f"{len(ORDER)} references cited, numbered in order of first appearance.")
